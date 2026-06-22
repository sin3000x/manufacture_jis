from __future__ import annotations

from collections.abc import Sequence
from pathlib import Path
import unicodedata

import pandas as pd


def write_df_to_excel(
    df: pd.DataFrame,
    path: str | Path,
    merge_dict: dict[tuple[str, ...], list[str]] | None = None,
) -> None:
    """将 DataFrame 写入 xlsx，并按 merge_dict 做单元格合并+居中。

    merge_dict: key 为分组依据列（tuple），value 为在该分组内需要合并的列列表。
    df 应已按 merge_dict 的 key 列排好序。
    """
    path = Path(path)
    df.to_excel(path, index=False)

    if not merge_dict:
        return

    from openpyxl import load_workbook
    from openpyxl.styles import Alignment

    wb = load_workbook(path)
    ws = wb.active
    col_index = {name: i + 1 for i, name in enumerate(df.columns)}
    header_row = 1

    center = Alignment(horizontal="center", vertical="center")

    normalized_merge_dict = _normalize_merge_dict(merge_dict)
    ordered_keys = sorted(normalized_merge_dict, key=len)

    _apply_auto_width(ws, df)

    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = center

    for group_keys in ordered_keys:
        target_cols = normalized_merge_dict[group_keys]
        group_cols = [col for col in group_keys if col in col_index]
        target_col_indices = [col_index[c] for c in target_cols if c in col_index]
        if not group_cols or not target_col_indices:
            continue

        for _, group in df.groupby(group_cols, sort=False):
            rows = group.index.tolist()
            if not rows:
                continue

            start_excel = rows[0] + header_row + 1
            end_excel = rows[-1] + header_row + 1
            for ci in target_col_indices:
                if start_excel != end_excel:
                    ws.merge_cells(
                        start_row=start_excel,
                        end_row=end_excel,
                        start_column=ci,
                        end_column=ci,
                    )
                ws.cell(row=start_excel, column=ci).alignment = center

    wb.save(path)


def _normalize_merge_dict(
    merge_dict: dict[tuple[str, ...] | Sequence[str], list[str]],
) -> dict[tuple[str, ...], list[str]]:
    normalized: dict[tuple[str, ...], list[str]] = {}
    for raw_group_keys, target_cols in merge_dict.items():
        group_keys = tuple(raw_group_keys)
        dedup_targets: list[str] = []
        for col in target_cols:
            if col not in dedup_targets:
                dedup_targets.append(col)
        normalized[group_keys] = dedup_targets
    return normalized


def _apply_auto_width(ws, df: pd.DataFrame) -> None:
    for column_cells in ws.columns:
        first_cell = column_cells[0]
        col_letter = first_cell.column_letter
        header_value = first_cell.value
        values = [header_value]
        for cell in column_cells[1:]:
            values.append(cell.value)

        width = 0
        for value in values:
            width = max(width, _display_width(value))

        if width <= 0:
            continue

        ws.column_dimensions[col_letter].width = min(width + 2, 40)


def _display_width(value) -> int:
    if value is None:
        return 0
    text = str(value)
    width = 0
    for char in text:
        if unicodedata.east_asian_width(char) in {"F", "W", "A"}:
            width += 2
        else:
            width += 1
    return width
