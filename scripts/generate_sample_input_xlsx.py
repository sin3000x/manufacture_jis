"""Generate sample XLSX input data for the manufacture JIS scheduler.

The generated workbook follows the input sheet descriptions in ``doc.txt``:

- 01排产信息（输入）
- 02输入（SR拆分）
- 03包规基表
- 04供应商作息&车规基表

Usage:
    uv run python scripts/generate_sample_input_xlsx.py
    uv run python scripts/generate_sample_input_xlsx.py --output data/custom_input.xlsx
"""

from __future__ import annotations

import argparse
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

DEFAULT_OUTPUT_PATH = Path("data/sample_input.xlsx")

SHEETS = {
    "production_plan": "01排产信息（输入）",
    "split_result": "02输入（SR拆分）",
    "package_rules": "03包规基表",
    "truck_rules": "04供应商作息&车规基表",
}

PRODUCTION_PLAN_HEADERS = [
    "线体",
    "任务令",
    "原材料编码",
    "计划量",
    "休息时间",
    "开工时间",
    "完工时间",
]

PRODUCTION_PLAN_ROWS = [
    [
        "line1",
        "mfg_order1",
        "item1",
        120,
        2.833,
        "2026-06-10 00:39:25",
        "2026-06-10 11:30:01",
    ],
]

SPLIT_RESULT_HEADERS = ["物料编码", "供应商名称", "货位", "需求数量", "需求ID"]
SPLIT_RESULT_ROWS = [
    ["item1", "supplier1", "loc1", 60, "demand1"],
    ["item1", "supplier1", "loc1", 60, "demand2"],
]

PACKAGE_RULE_HEADERS = ["编码", "包规"]
PACKAGE_RULE_ROWS = [
    ["item1", 5],
]

TRUCK_RULE_HEADERS = ["供应商名称", "车规（板/车）"]
TRUCK_RULE_ROWS = [
    ["supplier1", 18],
    ["supplier2", 22],
    ["supplier3", 18],
]


def write_sheet(workbook: Workbook, title: str, headers: list[str], rows: list[list[object]]) -> None:
    """Create a worksheet with styled headers and auto-sized columns."""
    worksheet = workbook.create_sheet(title=title)
    worksheet.append(headers)

    for row in rows:
        worksheet.append(row)

    header_fill = PatternFill(fill_type="solid", fgColor="D9EAF7")
    header_font = Font(bold=True)
    center_alignment = Alignment(horizontal="center", vertical="center")

    for cell in worksheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center_alignment

    for row in worksheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="center")

    for column_cells in worksheet.columns:
        max_length = max(len(str(cell.value)) if cell.value is not None else 0 for cell in column_cells)
        column_letter = get_column_letter(column_cells[0].column)
        worksheet.column_dimensions[column_letter].width = min(max(max_length + 2, 12), 24)

    worksheet.freeze_panes = "A2"


def generate_workbook(output_path: Path) -> Path:
    """Generate the sample input workbook and return its path."""
    output_path.parent.mkdir(parents=True, exist_ok=True)

    workbook = Workbook()
    default_sheet = workbook.active
    workbook.remove(default_sheet)

    write_sheet(
        workbook,
        SHEETS["production_plan"],
        PRODUCTION_PLAN_HEADERS,
        PRODUCTION_PLAN_ROWS,
    )
    write_sheet(
        workbook,
        SHEETS["split_result"],
        SPLIT_RESULT_HEADERS,
        SPLIT_RESULT_ROWS,
    )
    write_sheet(
        workbook,
        SHEETS["package_rules"],
        PACKAGE_RULE_HEADERS,
        PACKAGE_RULE_ROWS,
    )
    write_sheet(
        workbook,
        SHEETS["truck_rules"],
        TRUCK_RULE_HEADERS,
        TRUCK_RULE_ROWS,
    )

    workbook.save(output_path)
    return output_path


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Generate sample input XLSX for manufacture JIS scheduling.")
    parser.add_argument(
        "--output",
        type=Path,
        default=DEFAULT_OUTPUT_PATH,
        help=f"Output XLSX path. Defaults to {DEFAULT_OUTPUT_PATH}",
    )
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    output_path = generate_workbook(args.output)
    print(f"Generated input workbook: {output_path}")


if __name__ == "__main__":
    main()
