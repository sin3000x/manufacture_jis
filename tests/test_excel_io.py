from __future__ import annotations

from io import BytesIO
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

from manufacture_jis import schedule_dataframe_from_xlsx

OUTPUT_COLUMNS = [
    "供应商",
    "车次",
    "物料编码",
    "任务令",
    "线体",
    "货位",
    "需求ID",
    "满足数量",
    "包规",
    "板数",
    "总板数",
    "车规",
    "装载率",
    "计划到货时间",
]


def test_schedule_dataframe_from_xlsx_accepts_path_and_bytesio(tmp_path: Path) -> None:
    input_path = Path("data/sample_input.xlsx")

    df_from_path = schedule_dataframe_from_xlsx(input_path)
    assert list(df_from_path.columns) == OUTPUT_COLUMNS
    assert not df_from_path.empty
    assert set(df_from_path["供应商"]) == {"supplier1"}
    assert set(df_from_path["车次"]) == {"supplier1-T001", "supplier1-T002"}
    assert df_from_path["满足数量"].sum() == 120
    assert df_from_path.groupby("车次")["总板数"].first().le(df_from_path.groupby("车次")["车规"].first()).all()
    assert all(ts.hour not in {1, 2, 7, 12, 13, 18} for ts in df_from_path["计划到货时间"])

    xlsx_bytes = BytesIO(input_path.read_bytes())
    df_from_bytes = schedule_dataframe_from_xlsx(xlsx_bytes)
    pd.testing.assert_frame_equal(df_from_path, df_from_bytes)

    output_path = tmp_path / "result.xlsx"
    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        df_from_path.to_excel(writer, sheet_name="结果表", index=False)

    written = load_workbook(output_path, data_only=True)["结果表"]
    assert written.max_row == len(df_from_path) + 1
    assert [cell.value for cell in written[1]] == OUTPUT_COLUMNS
    assert written[2][0].value == "supplier1"
    assert written[2][1].value == "supplier1-T001"
